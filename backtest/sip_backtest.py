#!/usr/bin/env python3
"""
SIP backtest: ₹20,00,000 initial + ₹20,000/month, on the config
  breakout entry + Nifty 500 6-month RS + 20% ratcheting trailing stop + 4% sizing.

Outputs:
  backtest/results/sip_equity.svg   — portfolio value vs total invested (log)
  backtest/results/CW2sigma_SIP_backtest.xlsx — Summary + Equity Curve + Trades sheets
Prints headline numbers (XIRR, total invested, final value, max drawdown).
In-sample, survivorship-biased.
"""
import os, sys, datetime as dt
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import cw2sigma_backtest as bt
import calmar_experiments as ce
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

INIT, SIP = 2_000_000, 20_000

def xirr(cfs):
    d0 = min(d for d, _ in cfs)
    npv = lambda r: sum(cf / (1.0 + r) ** ((d - d0).days / 365.25) for d, cf in cfs)
    lo, hi = -0.9999, 100.0
    flo = npv(lo)
    for _ in range(300):
        mid = (lo + hi) / 2.0; fm = npv(mid)
        if abs(fm) < 1.0:
            return mid
        if (fm > 0) == (flo > 0):
            lo, flo = mid, fm
        else:
            hi = mid
    return (lo + hi) / 2.0

def maxdd(curve):
    peak = -1e18; mdd = 0.0
    for _, v in curve:
        peak = max(peak, v)
        mdd = max(mdd, (peak - v) / peak if peak > 0 else 0.0)
    return mdd

def main():
    symbols, data, master = ce.setup()
    bt.annotate_rs(data, symbols, lookback=26, index_sym="%5ECRSLDX")
    bt.POS_PCT = 0.04
    r = bt.simulate(symbols, data, master, bt.EXIT_MODES["ATR trail only"],
                    trail_type="pct", pct_trail=0.20, rs_entry=True, size_mode="fixed",
                    max_pos=50, monthly_add=SIP)
    bt.POS_PCT = 0.02
    curve, trades, contribs = r["curve"], r["trades"], r["contributions"]

    # invested-capital line (lump + cumulative SIP) aligned to the weekly curve
    cum, ci, invested_at = INIT, 0, []
    contribs_sorted = sorted(contribs)
    for d, _ in curve:
        while ci < len(contribs_sorted) and contribs_sorted[ci][0] <= d:
            cum += contribs_sorted[ci][1]; ci += 1
        invested_at.append(cum)

    total_inv = INIT + sum(a for _, a in contribs)
    final = curve[-1][1]
    yrs = (master[-1] - master[0]).days / 365.25
    cfs = [(master[0], -INIT)] + [(d, -a) for d, a in contribs] + [(curve[-1][0], final)]
    x = xirr(cfs)
    mdd = maxdd(curve)
    wins = [t for t in trades if t["ret"] > 0]
    print(f"Window {master[0]} -> {master[-1]} ({yrs:.1f} yrs)")
    print(f"Contributions: {len(contribs)} months x ₹{SIP:,}  +  ₹{INIT:,} lump")
    print(f"Total invested: ₹{total_inv:,.0f}")
    print(f"Final value:    ₹{final:,.0f}   ({final/total_inv:.2f}x invested)")
    print(f"XIRR:           {x*100:.1f}%")
    print(f"Max drawdown:   {mdd*100:.1f}%")
    print(f"Trades:         {len(trades)}   win {len(wins)/len(trades)*100:.0f}%")

    # ---- equity curve chart: portfolio value vs total invested ----
    series = [(f"Portfolio value (XIRR {x*100:.1f}%)", "#2563eb", curve),
              ("Total invested (lump + SIP)", "#9ca3af", list(zip([d for d, _ in curve], invested_at)))]
    bt.make_svg(series, os.path.join(bt.OUTDIR, "results", "sip_equity.svg"),
                "Rs 20L + Rs 20k/month — portfolio value vs invested (log)")

    # ---- Excel workbook ----
    wb = openpyxl.Workbook()
    ARIAL = "Arial"
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(name=ARIAL, bold=True, color="FFFFFF")
    bold = Font(name=ARIAL, bold=True)
    base = Font(name=ARIAL)

    def style_header(ws, ncol):
        for c in range(1, ncol + 1):
            cell = ws.cell(1, c); cell.fill = hdr_fill; cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"

    # Summary
    s = wb.active; s.title = "Summary"
    rows = [
        ("CW 2σ SIP backtest", ""),
        ("Rules", "Upper Bollinger breakout + Nifty 500 6-month outperformance + 20% ratcheting trailing stop + 4% sizing (weekly)"),
        ("Window", f"{master[0]} to {master[-1]}  ({yrs:.1f} years)"),
        ("Initial investment (₹)", INIT),
        ("Monthly addition (₹)", SIP),
        ("Number of monthly additions", len(contribs)),
        ("Total invested (₹)", total_inv),
        ("Final portfolio value (₹)", final),
        ("Total profit (₹)", final - total_inv),
        ("Growth multiple (x invested)", round(final / total_inv, 2)),
        ("XIRR (money-weighted annual return)", x),
        ("Max drawdown", -mdd),
        ("Total trades", len(trades)),
        ("Win rate", len(wins) / len(trades)),
        ("Note", "In-sample, survivorship-biased (today's Nifty 500). Educational, not advice."),
    ]
    for i, (k, v) in enumerate(rows, 1):
        s.cell(i, 1, k).font = bold if i == 1 else base
        c = s.cell(i, 2, v); c.font = base
        if isinstance(v, (int, float)) and k not in ("Number of monthly additions", "Total trades"):
            if "₹" in k:
                c.number_format = '#,##0'
            elif k in ("XIRR (money-weighted annual return)", "Max drawdown", "Win rate"):
                c.number_format = '0.0%'
    s.column_dimensions["A"].width = 34; s.column_dimensions["B"].width = 90

    # Equity Curve
    ec = wb.create_sheet("Equity Curve")
    ec.append(["Date", "Portfolio Value (₹)", "Total Invested (₹)", "Drawdown %"])
    peak = curve[0][1]
    for (d, v), inv in zip(curve, invested_at):
        peak = max(peak, v)
        dd = -(peak - v) / peak if peak > 0 else 0
        ec.append([d, round(v, 0), inv, dd])
    for row in ec.iter_rows(min_row=2, max_row=ec.max_row):
        row[0].number_format = "yyyy-mm-dd"; row[1].number_format = "#,##0"
        row[2].number_format = "#,##0"; row[3].number_format = "0.0%"
        for cc in row: cc.font = base
    style_header(ec, 4)
    for col, w in zip("ABCD", (12, 20, 18, 12)):
        ec.column_dimensions[col].width = w

    # Trades
    tr = wb.create_sheet("Trades")
    tr.append(["#", "Symbol", "Entry Date", "Entry ₹", "Qty", "Cost ₹",
               "Exit Date", "Exit ₹", "Return %", "P&L ₹", "Weeks Held"])
    for i, t in enumerate(sorted(trades, key=lambda z: z["entry_date"]), 1):
        tr.append([i, t["sym"].replace(".NS", "").replace("%26", "&"), t["entry_date"], round(t["entry"], 2),
                   t["shares"], round(t["cost"], 0), t["exit_date"], round(t["exit"], 2),
                   t["ret"], round(t["pnl"], 0), t["weeks"]])
    for row in tr.iter_rows(min_row=2, max_row=tr.max_row):
        row[2].number_format = "yyyy-mm-dd"; row[6].number_format = "yyyy-mm-dd"
        row[3].number_format = "#,##0.00"; row[5].number_format = "#,##0"
        row[7].number_format = "#,##0.00"; row[8].number_format = "0.0%"; row[9].number_format = "#,##0"
        for cc in row: cc.font = base
        row[8].font = Font(name=ARIAL, color="1B7F3B" if row[8].value and row[8].value > 0 else "B4231F")
    style_header(tr, 11)
    for col, w in zip("ABCDEFGHIJK", (5, 14, 12, 10, 8, 12, 12, 10, 10, 12, 11)):
        tr.column_dimensions[col].width = w

    out = os.path.join(bt.OUTDIR, "results", "CW2sigma_SIP_backtest.xlsx")
    wb.save(out)
    print(f"\nWrote {out}\nWrote backtest/results/sip_equity.svg")

if __name__ == "__main__":
    main()
